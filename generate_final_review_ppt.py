import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 Widescreen layout
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    # Theme Colors
    COLOR_PRIMARY = RGBColor(11, 25, 44)       # Deep Navy
    COLOR_ACCENT = RGBColor(0, 141, 218)       # Vibrant Tech Blue
    COLOR_BG = RGBColor(248, 249, 250)         # Soft Off-White
    COLOR_CARD = RGBColor(255, 255, 255)       # White Card
    COLOR_TEXT_DARK = RGBColor(33, 37, 41)     # Dark Slate
    COLOR_TEXT_MUTED = RGBColor(108, 117, 125) # Cool Gray

    def add_header(slide, title_text, category_text="BIG DATA ANALYTICS | FINAL PROJECT REVIEW"):
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.5), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_ACCENT

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.5), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_PRIMARY

    def add_card(slide, left, top, width, height, bg_color=COLOR_CARD, border_color=RGBColor(222, 226, 230)):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
        return shape

    # =========================================================================
    # SLIDE 1: Title Slide (Title, team members, register numbers, Faculty name)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = COLOR_PRIMARY
    bg1.line.fill.background()

    add_card(s1, Inches(1.2), Inches(1.0), Inches(10.9), Inches(5.5), COLOR_CARD)
    
    tb1 = s1.shapes.add_textbox(Inches(1.6), Inches(1.4), Inches(10.1), Inches(4.8))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "COURSE: BIG DATA ANALYTICS (23AID302)"
    p0.font.size = Pt(12)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_ACCENT

    p1 = tf1.add_paragraph()
    p1.text = "E-Commerce Review Spam Detection\nUsing Big Data"
    p1.font.size = Pt(30)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_PRIMARY
    p1.space_before = Pt(8)
    p1.space_after = Pt(16)

    p2 = tf1.add_paragraph()
    p2.text = "A Scalable Distributed Pipeline Using Hadoop HDFS, MapReduce, Scala, Apache Spark & Spark MLlib"
    p2.font.size = Pt(14)
    p2.font.italic = True
    p2.font.color.rgb = COLOR_TEXT_MUTED
    p2.space_after = Pt(28)

    p3 = tf1.add_paragraph()
    p3.text = "PRESENTED BY:"
    p3.font.size = Pt(11)
    p3.font.bold = True
    p3.font.color.rgb = COLOR_ACCENT

    p4 = tf1.add_paragraph()
    p4.text = "• Team Member Name(s):  [Your Name / Team Members]\n• Register Number(s):       [Your Register Number(s)]\n• Faculty / Supervisor:      [Faculty Name]"
    p4.font.size = Pt(13)
    p4.font.color.rgb = COLOR_TEXT_DARK
    p4.space_before = Pt(4)

    # =========================================================================
    # SLIDE 2: Problem Statement, Objectives, Why Big Data is Needed
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Problem Statement, Objectives & Why Big Data is Needed")

    # Card 1: Problem Statement
    add_card(s2, Inches(0.8), Inches(1.6), Inches(3.6), Inches(5.2))
    tb2_1 = s2.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(3.2), Inches(4.8))
    tf2_1 = tb2_1.text_frame
    tf2_1.word_wrap = True
    p = tf2_1.paragraphs[0]
    p.text = "1. Problem Statement"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    bullets1 = [
        "Modern e-commerce platforms (Amazon) rely heavily on user reviews for consumer trust.",
        "Malicious sellers flood products with fake positive spam reviews to artificially inflate ratings.",
        "Competitors post fake 1-star attack reviews to sabotage rival products.",
        "Manual moderation fails completely due to the massive volume of incoming customer reviews."
    ]
    for b in bullets1:
        p = tf2_1.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_TEXT_DARK
        p.space_before = Pt(8)

    # Card 2: Objectives
    add_card(s2, Inches(4.8), Inches(1.6), Inches(3.6), Inches(5.2))
    tb2_2 = s2.shapes.add_textbox(Inches(5.0), Inches(1.8), Inches(3.2), Inches(4.8))
    tf2_2 = tb2_2.text_frame
    tf2_2.word_wrap = True
    p = tf2_2.paragraphs[0]
    p.text = "2. Project Objectives"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    bullets2 = [
        "Distributed Ingestion: Store 232+ million reviews reliably across Hadoop HDFS DataNodes.",
        "Batch Processing: Run MapReduce and Scala preprocessing jobs to clean data.",
        "Storage Optimization: Convert 140 GB raw JSON into an optimized 50 GB Parquet format.",
        "Multi-Feature MLlib: Train an automated NLP spam classifier using Spark MLlib.",
        "Interactive Dashboard: Deploy a live web UI for real-time review spam inference."
    ]
    for b in bullets2:
        p = tf2_2.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_TEXT_DARK
        p.space_before = Pt(8)

    # Card 3: Why Big Data is Needed
    add_card(s2, Inches(8.8), Inches(1.6), Inches(3.7), Inches(5.2))
    tb2_3 = s2.shapes.add_textbox(Inches(9.0), Inches(1.8), Inches(3.3), Inches(4.8))
    tf2_3 = tb2_3.text_frame
    tf2_3.word_wrap = True
    p = tf2_3.paragraphs[0]
    p.text = "3. Why Big Data is Needed"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    bullets3 = [
        "Dataset Scale: Over 140 GB uncompressed text with ~232.8 million records.",
        "Single-Node Limit: Pandas & MySQL crash with Out-Of-Memory (OOM) errors.",
        "I/O Bottlenecks: Repeatedly parsing 140 GB of raw text takes hours on a single machine.",
        "Solution: Distributed block storage (HDFS) + in-memory parallel computation (Spark DAG)."
    ]
    for b in bullets3:
        p = tf2_3.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_TEXT_DARK
        p.space_before = Pt(8)

    # =========================================================================
    # SLIDE 3: Dataset: Source, Size, The 5 V's of Big Data, Key Attributes
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Dataset Profile & The 5 V's of Big Data")

    # Left Column: Profile & 5 Vs
    add_card(s3, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    tb3_1 = s3.shapes.add_textbox(Inches(1.0), Inches(1.75), Inches(5.2), Inches(4.9))
    tf3_1 = tb3_1.text_frame
    tf3_1.word_wrap = True

    p = tf3_1.paragraphs[0]
    p.text = "Dataset Profile & The 5 V's"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    ds_bullets = [
        ("Dataset Source: ", "Amazon Customer Reviews (All_Amazon_Review.json.gz)."),
        ("Storage Footprint: ", "33.5 GB gzip -> 140+ GB JSON -> 49.8 GB Parquet in HDFS."),
        ("1. VOLUME: ", "232.8M rows; 161 GB active replicated storage on 2 DataNodes."),
        ("2. VARIETY: ", "Semi-structured JSON, unstructured text, columnar Parquet & CSV."),
        ("3. VELOCITY: ", "High-throughput streaming into Parquet; distributed in-memory DAG."),
        ("4. VERACITY: ", "Tackling fake reviews, bot spam, unverified buyers & rating noise."),
        ("5. VALUE: ", "Restoring consumer trust, preventing fraud & protecting sellers.")
    ]
    for prefix, body in ds_bullets:
        p = tf3_1.add_paragraph()
        r1 = p.add_run()
        r1.text = "• " + prefix
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = COLOR_ACCENT
        r2 = p.add_run()
        r2.text = body
        r2.font.size = Pt(11)
        r2.font.color.rgb = COLOR_TEXT_DARK
        p.space_before = Pt(4)

    # Right Column: Key Attributes Table
    add_card(s3, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    tb3_2 = s3.shapes.add_textbox(Inches(7.0), Inches(1.8), Inches(5.3), Inches(4.8))
    tf3_2 = tb3_2.text_frame
    tf3_2.word_wrap = True

    p = tf3_2.paragraphs[0]
    p.text = "Key Attributes Used in Detection"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    attrs = [
        ("verified (Boolean)", "Primary Spam Marker. True = confirmed purchase; False = unverified/high-risk review."),
        ("reviewText (String)", "Natural language review content analyzed by Spark MLlib for spam vocabulary."),
        ("overall (Double)", "Star rating (1.0 to 5.0). Used to detect extreme rating clusters (5-star or 1-star)."),
        ("helpful (Array[Int])", "[Upvotes, Total]. Real reviews gain upvotes; spam almost always receives 0 upvotes."),
        ("asin (String)", "Amazon Product ID. Identifies products under high-volume spam attack waves."),
        ("reviewTime (String)", "Publication timestamp. Exposes abnormal sudden temporal review spikes.")
    ]
    for name, desc in attrs:
        p = tf3_2.add_paragraph()
        r1 = p.add_run()
        r1.text = "✔ " + name + ": "
        r1.font.bold = True
        r1.font.size = Pt(12)
        r1.font.color.rgb = COLOR_PRIMARY
        r2 = p.add_run()
        r2.text = desc
        r2.font.size = Pt(11)
        r2.font.color.rgb = COLOR_TEXT_DARK
        p.space_before = Pt(6)

    # =========================================================================
    # SLIDE 4: System Architecture Diagram: Data Source, Spark Cluster, Scala Code, ML, Output
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "System Architecture: End-to-End Pipeline")

    flow_steps = [
        ("1. Raw Data Source", "Amazon Reviews JSON.gz\n33.5 GB (Compressed)\n140+ GB Uncompressed\n232 Million Records", COLOR_CARD),
        ("2. HDFS Storage", "Hadoop 3.2.1 Cluster\n• 1 NameNode (:9870)\n• 2 DataNodes (Repl: 2)\n• 128 MB Block Size\n• 656 Parquet Blocks", COLOR_CARD),
        ("3. Scala & Spark Engine", "Apache Spark 3.1.1\n• Scala Preprocessing\n• Zero-Shuffle Streaming\n• 50 GB Parquet Dataset\n• DAG Execution Engine", COLOR_CARD),
        ("4. MLlib Pipeline", "Spark MLlib NLP\n• Tokenizer & TF-IDF\n• VectorAssembler\n• Logistic Regression\n• Standalone Cluster", COLOR_CARD),
        ("5. Outputs & Dashboard", "Analytics & Frontend\n• 7 HDFS CSV Reports\n• Interactive Dashboard\n• Real-Time Inference\n• Port 5000 Live UI", COLOR_CARD)
    ]

    card_w = Inches(2.25)
    card_h = Inches(3.6)
    spacing = Inches(0.18)

    for i, (stitle, sdesc, sbg) in enumerate(flow_steps):
        cx = Inches(0.8) + i * (card_w + spacing)
        add_card(s4, cx, Inches(1.6), card_w, card_h, sbg)
        
        tb = s4.shapes.add_textbox(cx + Inches(0.1), Inches(1.75), card_w - Inches(0.2), card_h - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = stitle
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = COLOR_ACCENT
        p.space_after = Pt(8)

        p2 = tf.add_paragraph()
        p2.text = sdesc
        p2.font.size = Pt(11)
        p2.font.color.rgb = COLOR_TEXT_DARK

    add_card(s4, Inches(0.8), Inches(5.4), Inches(11.7), Inches(1.5), RGBColor(235, 245, 255), COLOR_ACCENT)
    tb_bot = s4.shapes.add_textbox(Inches(1.0), Inches(5.5), Inches(11.3), Inches(1.3))
    tf_bot = tb_bot.text_frame
    tf_bot.word_wrap = True
    p = tf_bot.paragraphs[0]
    p.text = "Key Architectural Highlights:"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    
    arch_notes = [
        "Distributed Isolation: Deployed on Docker virtual bridge network connecting NameNode, 2 DataNodes, and Spark Master/Worker.",
        "High Availability: Replication Factor = 2 ensures zero data loss even if one DataNode goes offline.",
        "Two-Tier Spam Detection: Combines metadata behavioral screening (HDFS/Spark) with content-based NLP classification (MLlib)."
    ]
    for an in arch_notes:
        p = tf_bot.add_paragraph()
        p.text = "• " + an
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_TEXT_DARK
        p.space_before = Pt(2)

    # =========================================================================
    # SLIDE 5: Spark Workflow: RDD / DataFrame Operations, DAG, Partitioning, Caching Choices
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Spark Workflow: DataFrames, DAG, Partitioning & Caching")

    grid_items = [
        ("1. DataFrame & RDD Operations", [
            "High-Level API: Used Spark SQL DataFrames for clean, structured queries (groupBy, count, filter, select).",
            "Under the Hood: Every DataFrame operation maps directly to resilient RDDs (MapPartitionsRDD, ShuffledRowRDD).",
            "Lazy Evaluation: Spark builds an execution plan and only executes when an Action (.collect(), .write.csv()) is called."
        ]),
        ("2. DAG Execution Engine", [
            "DAGScheduler: Automatically converts DataFrame operations into a Directed Acyclic Graph (DAG).",
            "Stages: Splits execution into Stages at shuffle boundaries (e.g., wide transformations like groupBy).",
            "Tasks: Dispatches individual tasks per partition across available CPU cores (local[4])."
        ]),
        ("3. Partitioning Strategy", [
            "HDFS Alignment: Input Parquet dataset is split into 372 to 656 partitions matching 128 MB HDFS block sizes.",
            "Parallel Task Allocation: Each partition is processed by a dedicated task thread without cross-node network lockup.",
            "Zero Disk Shuffle: Avoided unnecessary .repartition() calls during conversion to prevent temporary file bloat."
        ]),
        ("4. Critical Caching Choices", [
            "Why We AVOIDED df.cache(): Uncompressed, the dataset expands to ~150 GB in memory, exceeding physical RAM and causing OutOfMemoryError.",
            "Streaming from Parquet: Parquet's columnar format allows reading only required columns without loading full rows into memory.",
            "Predicate Pushdown: Filters (e.g., overall, verified) are pushed to the storage layer, reading only relevant data."
        ])
    ]

    for i, (gtitle, gbullets) in enumerate(grid_items):
        row = i // 2
        col = i % 2
        gx = Inches(0.8) + col * Inches(5.95)
        gy = Inches(1.6) + row * Inches(2.65)
        gw = Inches(5.75)
        gh = Inches(2.45)

        add_card(s5, gx, gy, gw, gh)
        tb = s5.shapes.add_textbox(gx + Inches(0.2), gy + Inches(0.15), gw - Inches(0.4), gh - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = gtitle
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = COLOR_PRIMARY
        p.space_after = Pt(4)

        for b in gbullets:
            p = tf.add_paragraph()
            p.text = "• " + b
            p.font.size = Pt(11)
            p.font.color.rgb = COLOR_TEXT_DARK
            p.space_before = Pt(3)

    # =========================================================================
    # SLIDE 6: Scala Implementation: Key Code Snippets, Project Structure
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Scala Implementation: Project Structure & Code Snippet")

    # Left: Structure and Concepts
    add_card(s6, Inches(0.8), Inches(1.6), Inches(4.5), Inches(5.2))
    tb6_1 = s6.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(4.1), Inches(4.8))
    tf6_1 = tb6_1.text_frame
    tf6_1.word_wrap = True

    p = tf6_1.paragraphs[0]
    p.text = "Scala Role & Project Structure"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    scala_points = [
        ("Project Directory: ", "F:\\hadoop-project\\scala\\"),
        ("Source Script: ", "AmazonPreprocessing.scala"),
        ("Native Integration: ", "Apache Spark is natively written in Scala; Scala provides superior type-safety and JVM bytecode performance."),
        ("Functional Programming: ", "Utilized immutable variables (val), higher-order functions, and pattern matching expressions (when / otherwise)."),
        ("HDFS Ingestion: ", "Reads raw sample data directly from hdfs://namenode:9000/project/test/ and outputs cleaned data to /project/cleaned/reviews_scala.")
    ]
    for prefix, body in scala_points:
        p = tf6_1.add_paragraph()
        r1 = p.add_run()
        r1.text = "• " + prefix
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = COLOR_ACCENT
        r2 = p.add_run()
        r2.text = body
        r2.font.size = Pt(11)
        r2.font.color.rgb = COLOR_TEXT_DARK
        p.space_before = Pt(6)

    # Right: Code Snippet Card
    add_card(s6, Inches(5.6), Inches(1.6), Inches(6.9), Inches(5.2), RGBColor(15, 23, 42), RGBColor(30, 41, 59))
    tb6_2 = s6.shapes.add_textbox(Inches(5.8), Inches(1.75), Inches(6.5), Inches(4.9))
    tf6_2 = tb6_2.text_frame
    tf6_2.word_wrap = True

    p = tf6_2.paragraphs[0]
    p.text = "AmazonPreprocessing.scala (Core Functional Snippet)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = RGBColor(56, 189, 248)
    p.space_after = Pt(8)

    code_text = (
        "// 1. Ingest raw review data from HDFS\n"
        "val reviews = spark.read.json(\"hdfs://namenode:9000/.../sample.json\")\n\n"
        "// 2. Functional cleaning: Remove duplicates and missing text\n"
        "val cleanedReviews = reviews.dropDuplicates()\n"
        "  .filter(col(\"overall\").between(1.0, 5.0))\n"
        "  .filter(col(\"reviewText\").isNotNull)\n"
        "  .filter(length(trim(col(\"reviewText\"))) > 0)\n\n"
        "// 3. Conditional Feature Engineering (Sentiment derivation)\n"
        "val labeledReviews = cleanedReviews.withColumn(\"sentiment\",\n"
        "  when(col(\"overall\") <= 2.0, \"Negative\")\n"
        "    .when(col(\"overall\") === 3.0, \"Neutral\")\n"
        "    .when(col(\"overall\") >= 4.0, \"Positive\")\n"
        "    .otherwise(\"Unknown\")\n"
        ")\n\n"
        "// 4. Save cleaned dataset back to HDFS\n"
        "labeledReviews.write.parquet(\"hdfs://namenode:9000/.../reviews_scala\")"
    )

    p_code = tf6_2.add_paragraph()
    p_code.text = code_text
    p_code.font.name = "Consolas"
    p_code.font.size = Pt(10)
    p_code.font.color.rgb = RGBColor(241, 245, 249)

    # =========================================================================
    # SLIDE 7: Spark ML / DL Pipeline: Features, Algorithms, Tuning
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Spark MLlib Pipeline: Multi-Feature Architecture & Tuning")

    ml_cards = [
        ("1. Feature Engineering", [
            "Text Features (NLP):",
            "  • Tokenizer: Splits review sentences into lowercase word tokens.",
            "  • StopWordsRemover: Eliminates noise ('the', 'is', 'and', 'at').",
            "  • HashingTF: Maps tokens into 5,000-dimensional term frequencies.",
            "  • IDF: Computes Inverse Document Frequency to weight distinctive words.",
            "Metadata Features (Behavioral):",
            "  • Star Rating (overall): Detects extreme 5-star / 1-star skew.",
            "  • Character Length (review_len): Catches short bot boilerplate text.",
            "  • Helpful Votes (helpful_votes): Identifies lack of community upvotes."
        ]),
        ("2. VectorAssembler", [
            "Multimodal Fusion:",
            "  • VectorAssembler concatenates the 5,000 text features with the 3 numerical metadata features.",
            "  • Result: A unified feature vector fed into the machine learning classifier.",
            "Composite Ground-Truth Labeling:",
            "  • SPAM (1.0): (verified == False) AND (overall == 5.0 OR 1.0) AND (review_len < 150 OR helpful == 0).",
            "  • AUTHENTIC (0.0): Confirmed verified purchases with organic length.",
            "  • Avoids naively treating all unverified reviews as spam."
        ]),
        ("3. Algorithm & Tuning", [
            "Classification Algorithm:",
            "  • Logistic Regression: Scalable linear classifier computing optimal decision boundary via convex optimization.",
            "Hyperparameter Tuning:",
            "  • maxIter = 15: Sufficient iterations for gradient convergence.",
            "  • regParam = 0.01: L2 Regularization to prevent overfitting on rare vocabulary.",
            "  • Data Partitioning: 80% Training (48,000) / 20% Testing (12,000).",
            "Deployment:",
            "  • Distributed training across Spark Standalone Cluster."
        ])
    ]

    for i, (mtitle, mbullets) in enumerate(ml_cards):
        mx = Inches(0.8) + i * Inches(3.95)
        my = Inches(1.6)
        mw = Inches(3.8)
        mh = Inches(5.2)

        add_card(s7, mx, my, mw, mh)
        tb = s7.shapes.add_textbox(mx + Inches(0.15), my + Inches(0.15), mw - Inches(0.3), mh - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = mtitle
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = COLOR_PRIMARY
        p.space_after = Pt(6)

        for b in mbullets:
            p = tf.add_paragraph()
            p.text = b
            p.font.size = Pt(11)
            p.font.color.rgb = COLOR_TEXT_DARK
            p.space_before = Pt(3)

    # =========================================================================
    # SLIDE 8: Results: Metrics Table, Model Comparison, Charts, Business Insight
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Results: Metrics Table, Confusion Matrix & Business Insights")

    # Left: Metrics & Comparison Table
    add_card(s8, Inches(0.8), Inches(1.6), Inches(5.8), Inches(5.2))
    tb8_1 = s8.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.4), Inches(4.8))
    tf8_1 = tb8_1.text_frame
    tf8_1.word_wrap = True

    p = tf8_1.paragraphs[0]
    p.text = "Model Evaluation & Comparison"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    p_sub = tf8_1.add_paragraph()
    p_sub.text = "Multi-Feature Model vs. Baseline Text-Only Model (12,000 Test Records):"
    p_sub.font.size = Pt(11)
    p_sub.font.color.rgb = COLOR_TEXT_MUTED
    p_sub.space_after = Pt(8)

    comp_data = [
        ("Metric", "Baseline (Text-Only)", "Multi-Feature Model (Ours)"),
        ("Accuracy", "81.03%", "88.64% (+7.6%)"),
        ("Precision (Spam)", "78.20%", "86.72% (+8.5%)"),
        ("Recall (Spam)", "74.50%", "84.35% (+9.8%)"),
        ("F1-Score", "76.30%", "85.52% (+9.2%)")
    ]
    for row_idx, r_data in enumerate(comp_data):
        p = tf8_1.add_paragraph()
        if row_idx == 0:
            p.text = f"{r_data[0]:<18} | {r_data[1]:<20} | {r_data[2]}"
            p.font.bold = True
            p.font.size = Pt(11)
            p.font.color.rgb = COLOR_ACCENT
        else:
            p.text = f"{r_data[0]:<18} | {r_data[1]:<20} | {r_data[2]}"
            p.font.size = Pt(11)
            p.font.color.rgb = COLOR_TEXT_DARK
        p.space_before = Pt(4)

    p_cm_title = tf8_1.add_paragraph()
    p_cm_title.text = "\nConfusion Matrix (Test Set):"
    p_cm_title.font.bold = True
    p_cm_title.font.size = Pt(12)
    p_cm_title.font.color.rgb = COLOR_PRIMARY

    cm_lines = [
        "• True Negatives (TN - Authentic Approved): 6,612",
        "• True Positives (TP - Spam Caught):            4,025",
        "• False Positives (FP - False Alarm):              616",
        "• False Negatives (FN - Missed Spam):          747"
    ]
    for cml in cm_lines:
        p = tf8_1.add_paragraph()
        p.text = cml
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_TEXT_DARK

    # Right: Business Insights from HDFS 7 Reports
    add_card(s8, Inches(7.0), Inches(1.6), Inches(5.5), Inches(5.2))
    tb8_2 = s8.shapes.add_textbox(Inches(7.2), Inches(1.8), Inches(5.1), Inches(4.8))
    tf8_2 = tb8_2.text_frame
    tf8_2.word_wrap = True

    p = tf8_2.paragraphs[0]
    p.text = "Key Business Insights (232M Reviews)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    insights = [
        ("Extreme 5-Star Bias: ", "148.5M reviews (63.8%) are 5-star ratings. E-commerce platforms suffer from massive positive skew, requiring spam filtering to restore rating integrity."),
        ("Short Review Bot Clustering: ", "Reviews under 80 characters exhibit an 82.4% spam probability ('Great item', 'Loved it', 'Five stars')."),
        ("Helpfulness Validation: ", "Over 94% of confirmed spam reviews received 0 helpful community votes."),
        ("Product Targeting: ", "Top-selling electronic accessories (Fire Stick, Echo Dot, MicroSD cards) receive over 300,000+ reviews each, making them prime targets for bot review campaigns.")
    ]
    for prefix, body in insights:
        p = tf8_2.add_paragraph()
        r1 = p.add_run()
        r1.text = "✔ " + prefix
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = COLOR_PRIMARY
        r2 = p.add_run()
        r2.text = body
        r2.font.size = Pt(11)
        r2.font.color.rgb = COLOR_TEXT_DARK
        p.space_before = Pt(8)

    # =========================================================================
    # SLIDE 9: Challenges, Limitations, Conclusion and Future Work
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Engineering Challenges, Conclusion & Future Work")

    add_card(s9, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    tb9_1 = s9.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.2), Inches(4.8))
    tf9_1 = tb9_1.text_frame
    tf9_1.word_wrap = True

    p = tf9_1.paragraphs[0]
    p.text = "Engineering Challenges Overcome"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    challenges = [
        ("230 GB Virtual Disk Spill: ", "Initial conversion using .repartition(200) filled the virtual disk with shuffle spill files and crashed Docker. Resolved by zero-shuffle in-memory streaming directly to Parquet."),
        ("Memory Heap Exhaustion (OOM): ", "Caching 140 GB uncompressed text exceeded RAM. Solved by streaming directly from Parquet columnar blocks with predicate pushdown in local[4] mode."),
        ("HDFS Under-Replicated Blocks: ", "Default HDFS replication factor was 3 on 2 DataNodes. Resolved by setting cluster replication strictly to 2 (0 under-replicated blocks)."),
        ("Current Limitations: ", "Model evaluated primarily on English text reviews; ground truth utilizes composite behavioral heuristics as a proxy for physical purchase verification.")
    ]
    for prefix, body in challenges:
        p = tf9_1.add_paragraph()
        r1 = p.add_run()
        r1.text = "• " + prefix
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = COLOR_ACCENT
        r2 = p.add_run()
        r2.text = body
        r2.font.size = Pt(11)
        r2.font.color.rgb = COLOR_TEXT_DARK
        p.space_before = Pt(6)

    add_card(s9, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    tb9_2 = s9.shapes.add_textbox(Inches(7.0), Inches(1.8), Inches(5.3), Inches(4.8))
    tf9_2 = tb9_2.text_frame
    tf9_2.word_wrap = True

    p = tf9_2.paragraphs[0]
    p.text = "Conclusion & Future Work"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    concl_points = [
        ("Complete Lifecycle Mastery: ", "Successfully designed and deployed distributed storage (HDFS), batch processing (MapReduce), functional programming (Scala), query optimization (Parquet/DAG), and machine learning (Spark MLlib)."),
        ("Proven Scalability: ", "Successfully handled 232 Million records (140+ GB) on consumer hardware without memory collapse."),
        ("High-Performance Accuracy: ", "Multi-feature classifier achieved 88.64% accuracy and 85.52% F1-score with low false alarms."),
        ("Future Scope - Deep Learning: ", "Incorporate transformer architectures (BERT / DistilBERT) for nuanced contextual sarcasm detection."),
        ("Future Scope - Graph Analytics: ", "Implement GraphX to detect collusive reviewer rings and coordinated spam-farm botnets.")
    ]
    for prefix, body in concl_points:
        p = tf9_2.add_paragraph()
        r1 = p.add_run()
        r1.text = "✔ " + prefix
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = COLOR_PRIMARY
        r2 = p.add_run()
        r2.text = body
        r2.font.size = Pt(11)
        r2.font.color.rgb = COLOR_TEXT_DARK
        p.space_before = Pt(6)

    # =========================================================================
    # SLIDE 10: References and GitHub Link (Upload the PPT and Report)
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "References & Demonstration Deliverables")

    add_card(s10, Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2))
    tb10 = s10.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(11.1), Inches(4.8))
    tf10 = tb10.text_frame
    tf10.word_wrap = True

    p = tf10.paragraphs[0]
    p.text = "References, Code Repository & Submitted Artifacts"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(12)

    ref_items = [
        ("1. Dataset Reference: ", "Jianmo Ni, Jiacheng Li, Julian McAuley. 'Justifying Recommendations using Distantly-Labeled Reviews and Fine-Grained Aspects.' EMNLP 2019 (Amazon Customer Reviews)."),
        ("2. Apache Spark Reference: ", "Matei Zaharia et al. 'Apache Spark: A Unified Engine for Big Data Processing.' Communications of the ACM, 2016."),
        ("3. Hadoop Architecture: ", "Tom White. 'Hadoop: The Definitive Guide - Storage and Analysis at Internet Scale.' O'Reilly Media, 4th Edition."),
        ("4. Review Spam Literature: ", "Nitin Jindal and Bing Liu. 'Opinion Spam and Analysis.' Proceedings of the 2008 International Conference on Web Search and Data Mining (WSDM)."),
        ("5. GitHub Repository Link: ", "https://github.com/[username]/ecommerce-review-spam-detection-bigdata (Upload repository for code, Scala scripts, and Docker configs)."),
        ("6. Submitted Report File: ", "F:\\hadoop-project\\Ecommerce_Review_Spam_Detection_Project_Report.docx (Complete 11-Chapter Project Documentation)."),
        ("7. Live Interactive Dashboard: ", "http://localhost:5000/index.html (Features real-time review spam inference and dynamic analytics charts).")
    ]
    for prefix, body in ref_items:
        p = tf10.add_paragraph()
        r1 = p.add_run()
        r1.text = prefix
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = COLOR_ACCENT
        r2 = p.add_run()
        r2.text = body
        r2.font.size = Pt(11)
        r2.font.color.rgb = COLOR_TEXT_DARK
        p.space_before = Pt(6)

    out_path = r"F:\hadoop-project\Amazon_Reviews_BigData_Presentation.pptx"
    try:
        prs.save(out_path)
        print(f"Presentation successfully generated at: {out_path}")
    except PermissionError:
        alt_path = r"F:\hadoop-project\Amazon_Reviews_BigData_Presentation_Updated.pptx"
        prs.save(alt_path)
        print(f"Original file was open in PowerPoint! Successfully saved updated presentation to: {alt_path}")

if __name__ == "__main__":
    create_deck()

